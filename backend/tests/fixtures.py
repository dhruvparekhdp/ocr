"""Builders for realistic test documents, generated at test time so no binary fixtures live in git."""

import io
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
]


def _font(size: int) -> ImageFont.FreeTypeFont:
    for candidate in FONT_CANDIDATES:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size)


def text_image(lines: list[str], size: int = 32, width: int = 1400) -> Image.Image:
    img = Image.new("RGB", (width, int(size * 1.9 * len(lines) + size * 2)), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((size, size + i * int(size * 1.9)), line, fill="black", font=_font(size))
    return img


INVOICE_LINES = ["ACME TRADING LLC", "Invoice No: INV-2026-0042", "Date: 2026-03-14", "TOTAL DUE EUR 29.75"]
RECEIPT_LINES = ["CORNER CAFE", "Receipt R-7781", "Cappuccino 4.50", "Paid by card 4.50"]


def _draw_image(c: canvas.Canvas, img: Image.Image, x: float, y: float, width: float) -> None:
    """Draw keeping the aspect ratio: stretched glyphs are not what real scans look like."""
    c.drawImage(ImageReader(img), x, y, width=width, height=width * img.height / img.width)


def text_pdf(path: Path) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica", 11)
    c.drawString(50, 800, "ACME TRADING LLC")
    c.drawString(50, 780, "Invoice No: INV-2026-0042")
    c.drawString(380, 780, "Date: 2026-03-14")
    c.drawString(50, 700, "Widget A")
    c.drawString(300, 700, "2")
    c.drawString(450, 700, "25.00")
    c.drawString(50, 660, "TOTAL DUE EUR 29.75")
    c.showPage()
    c.setPageRotation(90)
    c.setFont("Helvetica", 11)
    c.drawString(50, 560, "Rotated page annex with terms and conditions")
    c.showPage()
    c.save()
    return path


def scanned_pdf(path: Path) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    _draw_image(c, text_image(INVOICE_LINES), 30, 550, width=535)
    c.showPage()
    c.save()
    return path


def mixed_pdf(path: Path) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica", 11)
    c.drawString(50, 800, "Expense report for March 2026, employee Jane Doe")
    _draw_image(c, text_image(RECEIPT_LINES), 50, 450, width=450)
    c.showPage()
    c.save()
    return path


def searchable_scan_pdf(path: Path) -> Path:
    """An image with an invisible OCR text layer on top, as produced by scanners."""
    c = canvas.Canvas(str(path), pagesize=A4)
    _draw_image(c, text_image(RECEIPT_LINES), 50, 450, width=450)
    text = c.beginText(60, 535)  # inside the image, which spans y=450..549
    text.setTextRenderMode(3)
    text.setFont("Helvetica", 14)
    for line in RECEIPT_LINES:
        text.textLine(line)
    c.drawText(text)
    c.showPage()
    c.save()
    return path


def encrypted_pdf(path: Path) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4, encrypt="secret")
    c.drawString(50, 800, "Confidential")
    c.showPage()
    c.save()
    return path


def rotated_photo(path: Path) -> Path:
    """Pixels stored sideways with an EXIF orientation tag, like a phone camera."""
    upright = text_image(RECEIPT_LINES)
    stored = upright.transpose(Image.Transpose.ROTATE_90)
    exif = Image.Exif()
    exif[0x0112] = 6  # display = rotate stored pixels 90 degrees clockwise
    stored.save(path, format="JPEG", exif=exif, quality=95)
    return path


def multipage_tiff(path: Path) -> Path:
    first, second = text_image(INVOICE_LINES), text_image(RECEIPT_LINES)
    first.save(path, format="TIFF", save_all=True, append_images=[second])
    return path


def workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(["Statement for Globex GmbH"])
    ws.merge_cells("A1:D1")
    ws.append(["Date", "Description", "Amount", "Account"])
    ws.append([date(2026, 1, 5), "Office rent", 1200.5, "00123"])
    ws.append([date(2026, 1, 9), "Hosting | cloud", 89, "00456"])
    ws["C5"] = "=SUM(C3:C4)"
    ws["H40"].number_format = "0.00"  # formatting-only cell must not widen the sheet

    buf = io.BytesIO()
    text_image(["Approved by CFO", "Ref APP-5521"], size=28, width=700).save(buf, format="PNG")
    buf.seek(0)
    picture = XLImage(buf)
    ws.add_image(picture, "F2")

    hidden = wb.create_sheet("Notes")
    hidden.sheet_state = "hidden"
    hidden["B2"] = "internal"
    wb.save(path)
    return path


def semicolon_csv(path: Path) -> Path:
    path.write_text("id;name;amount\n007;Globex;1.234,50\n008;Initech;99\n", encoding="utf-8")
    return path
