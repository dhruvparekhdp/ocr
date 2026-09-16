import csv
import io
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image, UnidentifiedImageError

from docai.ingest.errors import ParseError
from docai.ingest.layout import build_lines
from docai.ingest.model import CellValue, ParsedDocument, Sheet, SheetImage, column_letter
from docai.ingest.ocr import ocr_image

MAX_CELLS_PER_SHEET = 100_000
MAX_SHEETS = 50


def _cell_value(value) -> CellValue:
    match value:
        case None | bool() | int() | float():
            return value
        case str():
            return value.strip() or None
        case datetime():
            return value.isoformat(sep=" ") if value.time() != time() else value.date().isoformat()
        case date() | time():
            return value.isoformat()
        case timedelta():
            return str(value)
        case Decimal():
            return float(value)
    return str(value)


def _trim(rows: list[list[CellValue]]) -> list[list[CellValue]]:
    """Drop trailing empty rows/columns; openpyxl's dimensions include cells that only carry formatting."""
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, v in enumerate(r) if v is not None), default=0) for r in rows), default=0)
    return [r[:width] + [None] * (width - len(r[:width])) for r in rows]


def _limit(rows: list[list[CellValue]]) -> tuple[list[list[CellValue]], bool]:
    width = max((len(r) for r in rows), default=0)
    if not width or len(rows) * width <= MAX_CELLS_PER_SHEET:
        return rows, False
    return rows[: max(1, MAX_CELLS_PER_SHEET // width)], True


def parse_xlsx(path: Path) -> ParsedDocument:
    try:
        # data_only=True returns the value Excel last calculated; formulas are loaded separately so a
        # workbook written by a library (never opened in Excel, so no cached values) still shows something.
        # File objects, not paths: openpyxl rejects paths without an .xlsx extension (storage uses hashes).
        with path.open("rb") as f:
            values_wb = load_workbook(f, data_only=True)
        with path.open("rb") as f:
            formulas_wb = load_workbook(f, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as e:
        raise ParseError(f"cannot open workbook: {e}") from None

    doc = ParsedDocument(kind="spreadsheet")
    if len(values_wb.worksheets) > MAX_SHEETS:
        raise ParseError(f"workbook has {len(values_wb.worksheets)} sheets; the limit is {MAX_SHEETS}")

    uncached = 0
    for ws in values_wb.worksheets:
        fws = formulas_wb[ws.title]
        rows = []
        for vrow, frow in zip(ws.iter_rows(), fws.iter_rows(), strict=False):
            row = []
            for vcell, fcell in zip(vrow, frow, strict=False):
                value = _cell_value(vcell.value)
                if value is None and isinstance(fcell.value, str) and fcell.value.startswith("="):
                    value = fcell.value
                    uncached += 1
                row.append(value)
            rows.append(row)
        rows = _trim(rows)
        total = len(rows)
        rows, truncated = _limit(rows)

        sheet = Sheet(
            name=ws.title,
            hidden=ws.sheet_state != "visible",
            n_rows=total,
            n_cols=len(rows[0]) if rows else 0,
            rows=rows,
            merged=[str(r) for r in ws.merged_cells.ranges],
            truncated=truncated,
        )

        # openpyxl exposes embedded pictures only through the private `_images` list (stable for years).
        for img in getattr(ws, "_images", []):
            marker = img.anchor._from
            anchor = f"{column_letter(marker.col)}{marker.row + 1}"
            try:
                picture = Image.open(io.BytesIO(img._data()))
            except (UnidentifiedImageError, OSError):
                doc.warnings.append(f"sheet {ws.title!r}: unreadable image at {anchor}")
                continue
            lines = build_lines(ocr_image(picture), page_number=0)
            for line in lines:
                line.id = f"s{ws.title}-{anchor}-{line.id.split('-', 1)[1]}"
            sheet.images.append(SheetImage(anchor=anchor, lines=lines))

        if charts := len(getattr(ws, "_charts", [])):
            doc.warnings.append(f"sheet {ws.title!r}: {charts} chart(s) not extracted")
        doc.sheets.append(sheet)

    if uncached:
        doc.warnings.append(f"{uncached} formula cell(s) have no cached value; showing the formula instead")
    return doc


def parse_csv(path: Path, name: str) -> ParsedDocument:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ParseError("CSV must be UTF-8 encoded") from None

    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    # Values stay strings: guessing types would corrupt IDs with leading zeros and locale-formatted numbers.
    rows: list[list[CellValue]] = [[v.strip() or None for v in r] for r in csv.reader(io.StringIO(text), dialect)]
    width = max((len(r) for r in rows), default=0)
    rows = _trim([r + [None] * (width - len(r)) for r in rows])
    total = len(rows)
    rows, truncated = _limit(rows)

    doc = ParsedDocument(kind="spreadsheet")
    doc.sheets.append(
        Sheet(
            name=name,
            n_rows=total,
            n_cols=len(rows[0]) if rows else 0,
            rows=rows,
            truncated=truncated,
        )
    )
    return doc
